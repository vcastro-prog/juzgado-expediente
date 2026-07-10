from io import BytesIO
from datetime import datetime, timedelta
import shutil

import pandas as pd
import streamlit as st

from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


# =========================
# SEGURIDAD / LOGIN
# =========================
# La contraseña se configura en Streamlit Cloud:
# App > Settings > Secrets
#
# Debe existir:
# APP_PASSWORD = "tu_contraseña"
#
# En local puedes crear:
# .streamlit/secrets.toml
# con la misma línea.
# Ese archivo NO debe subirse a GitHub.

try:
    PASSWORD_CORRECTA = st.secrets["APP_PASSWORD"]
except Exception:
    PASSWORD_CORRECTA = None

if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

if not st.session_state.autenticado:
    st.set_page_config(
        page_title="Acceso privado",
        page_icon="🔒",
        layout="centered",
    )

    st.title("🔒 Acceso privado")
    st.caption("Introduce la contraseña para acceder al control de expedientes.")

    password = st.text_input(
        "Contraseña",
        type="password",
    )

    if st.button("Entrar", type="primary"):
        if PASSWORD_CORRECTA is None:
            st.error(
                "No se ha configurado APP_PASSWORD en los Secrets de Streamlit."
            )
        elif password == PASSWORD_CORRECTA:
            st.session_state.autenticado = True
            st.rerun()
        else:
            st.error("Contraseña incorrecta")

    st.stop()

from parser_pdf import extraer_expedientes, extraer_diagnostico_pdf, PARSER_VERSION, PDFConVariosJuzgadosError
from version import APP_VERSION, BUILD_DATE, APP_CHANGELOG
from database import (
    DB_PATH,
    guardar_importacion,
    cargar_expedientes,
    cargar_cambios,
    cargar_importaciones,
    cargar_detalle_expediente,
    guardar_favorito_nota,
    resetear_base,
)


st.set_page_config(
    page_title="Control de expedientes V2",
    page_icon="⚖️",
    layout="wide",
)

st.markdown(
    f"""
    <div style="display:flex; align-items:baseline; gap:14px; margin-bottom:0.4rem;">
        <h1 style="margin:0;">⚖️ Control de procedimientos del juzgado</h1>
        <span style="font-size:0.95rem; color:gray; white-space:nowrap;">
            Versión {APP_VERSION} · Build {BUILD_DATE} · {PARSER_VERSION}
        </span>
    </div>
    """,
    unsafe_allow_html=True,
)



def descargar_excel(df_dict):
    """
    Genera un Excel con formato:
    - primera fila congelada;
    - filtros activados;
    - cabecera azul con texto blanco;
    - filas alternas blanco / azul claro;
    - colores operativos:
        · modificados: verde claro;
        · sin iniciar trámite: amarillo claro;
        · archivados: gris claro, si aparece una columna que lo indique;
    - bordes suaves;
    - ancho de columnas ajustado.
    """
    output = BytesIO()

    color_cabecera = "1F4E78"
    color_fila_alterna = "DDEBF7"
    color_blanco = "FFFFFF"
    color_borde = "BFBFBF"

    color_modificado = "C6EFCE"      # verde claro
    color_sin_tramite = "FFF2CC"     # amarillo claro
    color_archivado = "E7E6E6"       # gris claro

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for nombre, df in df_dict.items():
            hoja = nombre[:31]
            df.to_excel(writer, index=False, sheet_name=hoja)

            ws = writer.book[hoja]

            max_row = ws.max_row
            max_col = ws.max_column

            if max_row == 0 or max_col == 0:
                continue

            # Congelar primera fila.
            ws.freeze_panes = "A2"

            # Activar autofiltro.
            ws.auto_filter.ref = ws.dimensions

            # Formatos.
            header_fill = PatternFill("solid", fgColor=color_cabecera)
            header_font = Font(color="FFFFFF", bold=True)
            even_fill = PatternFill("solid", fgColor=color_fila_alterna)
            odd_fill = PatternFill("solid", fgColor=color_blanco)
            fill_modificado = PatternFill("solid", fgColor=color_modificado)
            fill_sin_tramite = PatternFill("solid", fgColor=color_sin_tramite)
            fill_archivado = PatternFill("solid", fgColor=color_archivado)

            thin_border = Border(
                left=Side(style="thin", color=color_borde),
                right=Side(style="thin", color=color_borde),
                top=Side(style="thin", color=color_borde),
                bottom=Side(style="thin", color=color_borde),
            )

            # Cabecera.
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = thin_border

            # Índices de columnas para colores operativos.
            headers = {
                str(ws.cell(row=1, column=col_idx).value): col_idx
                for col_idx in range(1, max_col + 1)
            }

            col_estado_cambios = headers.get("Estado cambios")
            col_estado_tramite = headers.get("Estado F. Últ. Trámite")
            col_archivado = headers.get("Archivado detectado")

            # Filas alternas, bordes y colores operativos.
            for row_idx in range(2, max_row + 1):
                fill = even_fill if row_idx % 2 == 0 else odd_fill

                estado_cambios = ""
                estado_tramite = ""
                archivado = ""

                if col_estado_cambios:
                    estado_cambios = str(ws.cell(row=row_idx, column=col_estado_cambios).value or "")

                if col_estado_tramite:
                    estado_tramite = str(ws.cell(row=row_idx, column=col_estado_tramite).value or "")

                if col_archivado:
                    archivado = str(ws.cell(row=row_idx, column=col_archivado).value or "")

                # Prioridad de color:
                # 1. Modificado
                # 2. Sin iniciar trámite
                # 3. Archivado
                if "Modificado" in estado_cambios:
                    fill = fill_modificado
                elif "Sin iniciar" in estado_tramite:
                    fill = fill_sin_tramite
                elif archivado in ("1", "True", "true", "Sí", "Si"):
                    fill = fill_archivado

                for col_idx in range(1, max_col + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.fill = fill
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Ajuste de columnas con límites razonables.
            for col_idx in range(1, max_col + 1):
                col_letter = get_column_letter(col_idx)
                max_length = 0

                for row_idx in range(1, min(max_row, 500) + 1):
                    value = ws.cell(row=row_idx, column=col_idx).value
                    if value is None:
                        continue
                    max_length = max(max_length, len(str(value)))

                ancho = min(max(max_length + 2, 10), 45)
                ws.column_dimensions[col_letter].width = ancho

            # Altura cabecera.
            ws.row_dimensions[1].height = 24

    return output.getvalue()

def leer_backup_sqlite():
    if DB_PATH.exists():
        return DB_PATH.read_bytes()
    return None


def nombre_backup_sqlite():
    return "expedientes_backup_" + datetime.now().strftime("%Y%m%d_%H%M") + ".sqlite"


def info_backup_sqlite():
    if not DB_PATH.exists():
        return None

    stat = DB_PATH.stat()
    return {
        "ruta": str(DB_PATH),
        "tamano_mb": round(stat.st_size / (1024 * 1024), 2),
        "modificado": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
    }


def restaurar_backup_sqlite(archivo_subido):
    DB_PATH.parent.mkdir(exist_ok=True)
    with open(DB_PATH, "wb") as f:
        f.write(archivo_subido.getbuffer())


def normalizar_texto(valor):
    return str(valor or "").lower().strip()


def filtrar_por_texto(df, texto):
    if not texto:
        return df
    patron = texto.lower().strip()
    mascara = df.apply(
        lambda row: patron in " ".join(row.astype(str)).lower(),
        axis=1,
    )
    return df[mascara]



def filtrar_por_patron_expediente(df, patron):
    """
    Permite filtrar expedientes por terminación del número, sin contar el año.

    Ejemplos:
    - *5       -> 0000005/2026, 0001235/2024, etc.
    - *45      -> 0002345/2026, 0000045/2025, etc.
    - *345     -> 0002345/2026, etc.
    - *5/2025  -> expedientes cuyo número termina en 5 y año 2025.
    - 0002345/2026 -> búsqueda parcial normal.
    """
    if not patron:
        return df

    patron = str(patron).strip()
    if not patron:
        return df

    if not patron.startswith("*"):
        return df[df["Nº Proced."].astype(str).str.contains(patron, case=False, na=False)]

    patron_limpio = patron[1:].strip()

    if "/" in patron_limpio:
        terminacion, anio = patron_limpio.split("/", 1)
        terminacion = terminacion.strip()
        anio = anio.strip()
    else:
        terminacion = patron_limpio.strip()
        anio = ""

    def coincide(valor):
        valor = str(valor or "").strip()
        if "/" not in valor:
            return False

        numero, anio_exp = valor.split("/", 1)
        numero = numero.strip()
        anio_exp = anio_exp.strip()

        if terminacion and not numero.endswith(terminacion):
            return False

        if anio and anio_exp != anio:
            return False

        return True

    return df[df["Nº Proced."].apply(coincide)]


def filtrar_por_fecha_cambio(df_cambios, dias):
    if df_cambios.empty or dias == "Todos":
        return df_cambios

    ahora = datetime.now()
    if dias == "Hoy":
        limite = ahora.replace(hour=0, minute=0, second=0, microsecond=0)
    elif dias == "Últimos 7 días":
        limite = ahora - timedelta(days=7)
    elif dias == "Últimos 30 días":
        limite = ahora - timedelta(days=30)
    else:
        return df_cambios

    tmp = df_cambios.copy()
    tmp["_fecha"] = pd.to_datetime(tmp["Fecha cambio"], errors="coerce")
    tmp = tmp[tmp["_fecha"] >= limite]
    return tmp.drop(columns=["_fecha"])


with st.sidebar:
    st.caption(f"Versión {APP_VERSION} · {BUILD_DATE} · {PARSER_VERSION}")

    if st.button("Cerrar sesión"):
        st.session_state.autenticado = False
        st.rerun()

    st.divider()
    st.header("Importar PDFs")

    pdfs = st.file_uploader(
        "Selecciona uno o varios PDFs",
        type=["pdf"],
        accept_multiple_files=True,
    )

    if pdfs:
        if st.button("Procesar PDFs", type="primary"):
            total_registros = 0
            total_cambios = 0
            total_pdfs = len(pdfs)
            resumen_importacion = []

            with st.spinner("Leyendo PDFs y actualizando base de datos..."):
                for pdf in pdfs:
                    try:
                        diagnostico = extraer_diagnostico_pdf(pdf)

                        if diagnostico.get("error"):
                            raise Exception(diagnostico["error"])

                        registros = extraer_expedientes(pdf)
                        cambios = guardar_importacion(
                            registros,
                            nombre_archivo=pdf.name,
                        )

                        total_registros += len(registros)
                        total_cambios += len(cambios)

                        resumen_importacion.append(
                            {
                                "Archivo": pdf.name,
                                "Tipo PDF": diagnostico.get("tipo_pdf", ""),
                                "Páginas": diagnostico.get("paginas", ""),
                                "Parser": diagnostico.get("parser_version", ""),
                                "Expedientes leídos": len(registros),
                                "Cambios/nuevos detectados": len(cambios),
                                "Estado": "Procesado",
                            }
                        )

                    except PDFConVariosJuzgadosError as e:
                        resumen_importacion.append(
                            {
                                "Archivo": pdf.name,
                                "Tipo PDF": "",
                                "Páginas": "",
                                "Parser": PARSER_VERSION,
                                "Expedientes leídos": 0,
                                "Cambios/nuevos detectados": 0,
                                "Estado": "Bloqueado: varios juzgados detectados",
                            }
                        )
                        st.error(
                            f"{pdf.name}: el PDF contiene varios juzgados u órganos. "
                            "No se ha importado ningún dato de este archivo. "
                            "Vuelve a cargarlo separado, un PDF por juzgado."
                        )
                        st.write("Órganos detectados:")
                        for organo in e.juzgados_detectados:
                            st.write(f"- {organo}")

                    except Exception as e:
                        resumen_importacion.append(
                            {
                                "Archivo": pdf.name,
                                "Tipo PDF": "",
                                "Páginas": "",
                                "Parser": PARSER_VERSION,
                                "Expedientes leídos": 0,
                                "Cambios/nuevos detectados": 0,
                                "Estado": f"Error: {e}",
                            }
                        )
                        st.error(f"Error procesando {pdf.name}: {e}")

            st.success(f"PDFs procesados: {total_pdfs}")
            st.info(f"Expedientes leídos: {total_registros}")
            st.info(f"Cambios/nuevos detectados: {total_cambios}")

            if resumen_importacion:
                st.subheader("Resumen de importación")
                st.dataframe(
                    pd.DataFrame(resumen_importacion),
                    use_container_width=True,
                )

            st.warning(
                "Importación terminada. Descarga ahora una copia SQLite desde "
                "la sección 'Backup / Restauración' de la barra lateral. "
                "Así no tendrás que volver a cargar PDFs antiguos si Streamlit reinicia la app."
            )

    st.divider()
    st.header("Backup / Restauración")
    st.caption(
        "Guarda una copia después de cada importación. "
        "Si Streamlit pierde la base local, restaura este archivo y no tendrás que volver a subir PDFs antiguos."
    )

    info_backup = info_backup_sqlite()
    backup = leer_backup_sqlite()

    if info_backup:
        st.success("Base local disponible")
        st.caption(f"Última modificación: {info_backup['modificado']}")
        st.caption(f"Tamaño: {info_backup['tamano_mb']} MB")

        st.download_button(
            "Descargar copia SQLite",
            data=backup,
            file_name=nombre_backup_sqlite(),
            mime="application/octet-stream",
            help="Guarda este archivo en tu PC, NAS o nube privada.",
        )
    else:
        st.info("Aún no hay base SQLite para descargar.")

    backup_subido = st.file_uploader(
        "Restaurar copia SQLite",
        type=["sqlite", "db"],
        help="Sube una copia SQLite previamente descargada.",
    )

    if backup_subido is not None:
        st.warning(
            "Restaurar una copia sustituirá la base local actual por el archivo subido."
        )
        if st.button("Restaurar copia SQLite", type="primary"):
            restaurar_backup_sqlite(backup_subido)
            st.success("Copia restaurada. Recarga la página para ver los datos.")

    st.divider()
    st.warning("Zona de mantenimiento")
    confirmar = st.checkbox("Confirmo que quiero borrar la base local")
    if confirmar and st.button("Borrar base de datos"):
        resetear_base()
        st.success("Base reiniciada. Recarga la página si no ves los cambios.")


df = cargar_expedientes()
df_cambios = cargar_cambios()
df_importaciones = cargar_importaciones()

if "filtro_juzgado_click" not in st.session_state:
    st.session_state.filtro_juzgado_click = ""

if "df_exportar_filtrado" not in st.session_state:
    st.session_state.df_exportar_filtrado = pd.DataFrame()
if "descripcion_exportacion" not in st.session_state:
    st.session_state.descripcion_exportacion = "Sin filtros aplicados"
if "columnas_exportar_visibles" not in st.session_state:
    st.session_state.columnas_exportar_visibles = []

with st.sidebar:
    st.divider()
    st.header("Datos generales")
    st.metric("Expedientes totales", len(df))
    st.metric("Favoritos", int(df["Favorito"].sum()) if not df.empty and "Favorito" in df else 0)
    st.metric("Cambios históricos", len(df_cambios))
    st.metric("Importaciones", len(df_importaciones))
    st.caption(
        "Última importación: "
        + (df_importaciones["Fecha importación"].iloc[0] if len(df_importaciones) else "Sin datos")
    )

tab_actuales, tab_modificados, tab_detalle, tab_cambios, tab_importaciones, tab_exportar = st.tabs(
    [
        "Expedientes actuales",
        "Modificados recientes",
        "Detalle expediente",
        "Histórico cambios",
        "Importaciones",
        "Exportar",
    ]
)


with tab_actuales:
    st.subheader("Expedientes actuales")

    if df.empty:
        st.session_state.df_exportar_filtrado = pd.DataFrame()
        st.session_state.descripcion_exportacion = "Sin expedientes"
        st.info("Aún no hay expedientes. Sube uno o varios PDFs desde la barra lateral.")
    else:
        with st.expander("Filtros avanzados", expanded=True):

            col_reset_1, col_reset_2 = st.columns([1, 5])

            with col_reset_1:
                if st.button("🔄 Resetear filtros"):
                    for key in list(st.session_state.keys()):
                        if key.startswith("filtro_"):
                            del st.session_state[key]
                    st.rerun()
            col1, col2, col3, col4 = st.columns(4)

            juzgados = sorted([x for x in df["Juzgado"].dropna().unique() if x])
            numeros_juzgado = sorted([x for x in df["Nº Juzgado"].dropna().unique() if x])
            tipos_organo = sorted([x for x in df["Tipo órgano"].dropna().unique() if x])
            secciones = sorted([x for x in df["Sección"].dropna().unique() if x])
            procedimientos = sorted([x for x in df["Procedimiento"].dropna().unique() if x])
            fases = sorted([x for x in df["Fase Procesal"].dropna().unique() if x])
            materias = sorted([x for x in df["Materia"].dropna().unique() if x])
            anios = sorted([x for x in df["Año"].dropna().unique() if x])
            anios_ultimo_tramite = sorted([x for x in df["Año Últ. Trámite"].dropna().unique() if x])
            tipos_resolucion = sorted([x for x in df["Tipo Resolución"].dropna().unique() if x])
            estados_resolucion = sorted([x for x in df["Estado Resolución"].dropna().unique() if x])
            intervenciones = sorted([x for x in df["Intervención"].dropna().unique() if x])
            anios_resolucion = sorted([x for x in df["Año Resolución"].dropna().unique() if x])

            filtro_numero = col1.text_input(
                "Nº procedimiento / patrón",
                key="filtro_numero_patron",
                help="Ejemplos: 0002345/2026, *5, *45, *345 o *5/2025. El patrón * busca por terminación del número sin contar el año.",
            )
            filtro_anio = col2.multiselect(
                "Año del procedimiento",
                anios,
                key="filtro_anio_procedimiento",
            )
            filtro_favoritos = col3.checkbox(
                "Solo favoritos",
                key="filtro_solo_favoritos",
            )
            filtro_archivados = col4.selectbox(
                "Estado general",
                ["Todos", "Excluir archivados", "Solo archivados"],
                key="filtro_estado_general",
            )

            filtro_cambios = st.selectbox(
                "Cambios recientes",
                [
                    "Todos",
                    "Solo modificados hoy",
                    "Modificados últimos 7 días",
                    "Modificados últimos 30 días",
                ],
                key="filtro_cambios_recientes",
            )

            st.caption(
                "Filtro de nº procedimiento: usa *5, *45 o *345 para buscar expedientes cuyo número termine así. "
                "Usa *5/2025 para limitar además al año 2025."
            )

            st.markdown("#### Filtros del Libro de Resoluciones")
            cr1, cr2, cr3 = st.columns(3)
            filtro_num_resolucion = cr1.text_input("Nº resolución", key="filtro_num_resolucion")
            filtro_anio_resolucion = cr2.multiselect("Año de resolución", anios_resolucion, key="filtro_anio_resolucion")
            filtro_tipo_resolucion = cr3.multiselect("Tipo de resolución", tipos_resolucion, key="filtro_tipo_resolucion")
            cr4, cr5, cr6 = st.columns(3)
            filtro_estado_resolucion = cr4.multiselect("Estado de resolución", estados_resolucion, key="filtro_estado_resolucion")
            filtro_intervencion = cr5.multiselect("Intervención", intervenciones, key="filtro_intervencion")
            filtro_interviniente = cr6.text_input("Interviniente contiene", key="filtro_interviniente")

            col_fecha1, col_fecha2 = st.columns(2)
            filtro_anio_ultimo_tramite = col_fecha1.multiselect(
                "Año de Fecha Últ. Trámite",
                anios_ultimo_tramite,
                help="Filtra por el año de la fecha del último trámite.",
            )
            filtro_fecha_vacia = col_fecha2.checkbox(
                "Solo Fecha Últ. Trámite vacía",
                help="Muestra expedientes sin fecha de último trámite. Sirve para localizar los que no han empezado a tramitarse.",
            )

            juzgado = st.multiselect("Juzgado / órgano", juzgados)
            col_j1, col_j2, col_j3 = st.columns(3)
            filtro_num_juzgado = col_j1.multiselect("Nº juzgado", numeros_juzgado)
            filtro_tipo_organo = col_j2.multiselect("Tipo órgano", tipos_organo)
            filtro_seccion = col_j3.multiselect("Sección", secciones)
            procedimiento = st.multiselect("Procedimiento", procedimientos)
            fase = st.multiselect("Fase procesal", fases)
            materia = st.multiselect("Materia", materias)

            if st.session_state.filtro_juzgado_click:
                st.info(
                    "Filtro rápido activo: "
                    + st.session_state.filtro_juzgado_click
                )
                if st.button("Quitar filtro rápido de juzgado"):
                    st.session_state.filtro_juzgado_click = ""
                    st.rerun()

            col5, col6 = st.columns(2)
            tramite_contiene = col5.text_input("Último trámite contiene")
            texto_general = col6.text_input("Búsqueda general")

        filtrado = df.copy()

        if filtro_numero:
            filtrado = filtrar_por_patron_expediente(filtrado, filtro_numero)

        if filtro_anio:
            filtrado = filtrado[filtrado["Año"].isin(filtro_anio)]

        if filtro_num_resolucion:
            filtrado = filtrado[filtrado["Nº Resolución"].astype(str).str.contains(filtro_num_resolucion, case=False, na=False)]
        if filtro_anio_resolucion:
            filtrado = filtrado[filtrado["Año Resolución"].isin(filtro_anio_resolucion)]
        if filtro_tipo_resolucion:
            filtrado = filtrado[filtrado["Tipo Resolución"].isin(filtro_tipo_resolucion)]
        if filtro_estado_resolucion:
            filtrado = filtrado[filtrado["Estado Resolución"].isin(filtro_estado_resolucion)]
        if filtro_intervencion:
            filtrado = filtrado[filtrado["Intervención"].isin(filtro_intervencion)]
        if filtro_interviniente:
            filtrado = filtrado[filtrado["Interviniente"].astype(str).str.contains(filtro_interviniente, case=False, na=False)]

        if filtro_favoritos:
            filtrado = filtrado[filtrado["Favorito"] == 1]

        if filtro_archivados == "Excluir archivados":
            filtrado = filtrado[filtrado["Archivado detectado"] == 0]
        elif filtro_archivados == "Solo archivados":
            filtrado = filtrado[filtrado["Archivado detectado"] == 1]

        if filtro_fecha_vacia:
            filtrado = filtrado[
                filtrado["Fecha Últ. Trámite"].isna()
                | (filtrado["Fecha Últ. Trámite"].astype(str).str.strip() == "")
            ]
        elif filtro_anio_ultimo_tramite:
            filtrado = filtrado[filtrado["Año Últ. Trámite"].isin(filtro_anio_ultimo_tramite)]

        if st.session_state.filtro_juzgado_click:
            filtrado = filtrado[
                filtrado["Órgano completo"].astype(str) == st.session_state.filtro_juzgado_click
            ]

        if juzgado:
            filtrado = filtrado[filtrado["Juzgado"].isin(juzgado)]

        if filtro_num_juzgado:
            filtrado = filtrado[filtrado["Nº Juzgado"].isin(filtro_num_juzgado)]

        if filtro_tipo_organo:
            filtrado = filtrado[filtrado["Tipo órgano"].isin(filtro_tipo_organo)]

        if filtro_seccion:
            filtrado = filtrado[filtrado["Sección"].isin(filtro_seccion)]

        if procedimiento:
            filtrado = filtrado[filtrado["Procedimiento"].isin(procedimiento)]
        if fase:
            filtrado = filtrado[filtrado["Fase Procesal"].isin(fase)]
        if materia:
            filtrado = filtrado[filtrado["Materia"].isin(materia)]

        if tramite_contiene:
            filtrado = filtrado[
                filtrado["Último trámite"].astype(str).str.contains(tramite_contiene, case=False, na=False)
            ]

        filtrado = filtrar_por_texto(filtrado, texto_general)

        st.subheader("Carga por juzgado del filtro aplicado")

        if filtrado.empty:
            st.info("No hay expedientes para los filtros seleccionados.")
        else:
            resumen_juzgado = (
                filtrado.groupby(
                    ["Nº Juzgado", "Tipo órgano", "Sección", "Órgano completo"],
                    dropna=False
                )
                .size()
                .reset_index(name="Expedientes")
                .sort_values("Expedientes", ascending=False)
            )
            total_filtrado = int(resumen_juzgado["Expedientes"].sum())
            resumen_juzgado["%"] = (
                resumen_juzgado["Expedientes"] / total_filtrado * 100
            ).round(1)

            # Mostramos TODOS los juzgados y permitimos pulsar para filtrar.
            # La tabla mantiene el orden por carga/porcentaje, pero las tarjetas se ordenan por nº de juzgado.
            resumen_cards = resumen_juzgado.copy()

            resumen_cards["_orden_juzgado"] = pd.to_numeric(
                resumen_cards["Nº Juzgado"],
                errors="coerce",
            )

            resumen_cards = (
                resumen_cards
                .sort_values(["_orden_juzgado", "Nº Juzgado"], na_position="last")
                .drop(columns=["_orden_juzgado"])
                .reset_index(drop=True)
            )

            for inicio in range(0, len(resumen_cards), 4):
                cols_juzgados = st.columns(min(4, len(resumen_cards) - inicio))

                for pos, (_, row) in enumerate(resumen_cards.iloc[inicio:inicio + 4].iterrows()):
                    organo = str(row["Órgano completo"])
                    etiqueta = "Juzgado " + str(row["Nº Juzgado"] or "sin nº")
                    with cols_juzgados[pos]:
                        st.metric(
                            label=etiqueta,
                            value=int(row["Expedientes"]),
                            delta=f'{row["%"]}% del filtro',
                        )
                        if st.button(
                            "Filtrar",
                            key="btn_filtrar_juzgado_" + str(inicio) + "_" + str(pos),
                            help="Mostrar solo expedientes de " + organo,
                        ):
                            st.session_state.filtro_juzgado_click = organo
                            st.rerun()

            st.dataframe(
                resumen_juzgado,
                use_container_width=True,
                hide_index=True,
            )

        st.caption(f"Expedientes mostrados en la tabla inferior: {len(filtrado)} de {len(df)} totales.")

        # La vista exacta para exportar se guarda después de seleccionar columnas visibles.

        columnas_disponibles = [
            "⭐", "Tipo documento", "Nº Resolución", "Año Resolución",
            "Fecha Dictado", "Hora Dictado", "Tipo Resolución",
            "Estado Resolución", "Intervención", "Interviniente",
            "Nº Proced.",
            "Año",
            "F. Aceptación",
            "Nº Juzgado",
            "Tipo órgano",
            "Sección",
            "Órgano completo",
            "Procedimiento",
            "Materia",
            "Fase Procesal",
            "Último trámite",
            "Fecha Últ. Trámite",
            "Año Últ. Trámite",
            "Estado F. Últ. Trámite",
            "Estado cambios",
            "Último cambio",
            "Nota",
            "Primera importación",
            "Última importación",
        ]

        columnas_por_defecto = [
            "⭐",
            "Nº Resolución",
            "Fecha Dictado",
            "Tipo Resolución",
            "Estado Resolución",
            "Nº Proced.",
            "Procedimiento",
            "Intervención",
            "Interviniente",
            "Año",
            "Nº Juzgado",
            "Materia",
            "Fase Procesal",
            "Último trámite",
            "Fecha Últ. Trámite",
            "Estado F. Últ. Trámite",
            "Estado cambios",
            "Último cambio",
            "Nota",
        ]

        st.subheader("Columnas visibles")
        columnas_visibles = st.multiselect(
            "Selecciona las columnas que quieres ver y exportar",
            options=columnas_disponibles,
            default=[c for c in columnas_por_defecto if c in columnas_disponibles],
            help="El Excel exportará exactamente estas columnas y los filtros aplicados.",
        )

        if not columnas_visibles:
            st.warning("Selecciona al menos una columna para mostrar la tabla.")
            columnas_visibles = ["Nº Proced."]

        # Evita nombres duplicados, ya que PyArrow/Streamlit no admite
        # DataFrames con columnas repetidas.
        columnas_visibles = list(dict.fromkeys(columnas_visibles))

        tabla = filtrado.copy()
        tabla["⭐"] = tabla["Favorito"].apply(lambda x: "⭐" if x else "")

        columnas_finales = [
            c for c in dict.fromkeys(columnas_visibles)
            if c in tabla.columns
        ]
        tabla_visible = tabla.loc[:, ~tabla.columns.duplicated()]
        tabla_visible = tabla_visible[columnas_finales].copy()

        st.session_state.df_exportar_filtrado = tabla_visible.copy()
        st.session_state.columnas_exportar_visibles = columnas_finales
        st.session_state.descripcion_exportacion = (
            f"Exportación filtrada: {len(tabla_visible)} de {len(df)} expedientes. "
            f"Columnas: {len(columnas_finales)}"
        )

        cambios_detectados = tabla_visible[
            tabla_visible["Estado cambios"].astype(str).str.contains("Modificado", na=False)
        ] if "Estado cambios" in tabla_visible.columns else pd.DataFrame()

        if not cambios_detectados.empty:
            st.success(
                f"Expedientes modificados detectados: {len(cambios_detectados)}"
            )

        st.dataframe(
            tabla_visible,
            use_container_width=True,
            height=650,
        )


with tab_modificados:
    st.subheader("Expedientes modificados recientemente")

    if df_cambios.empty:
        st.info("Aún no hay cambios registrados.")
    else:
        rango = st.radio(
            "Rango",
            ["Hoy", "Últimos 7 días", "Últimos 30 días", "Todos"],
            horizontal=True,
        )

        cambios_rango = filtrar_por_fecha_cambio(df_cambios, rango)

        if cambios_rango.empty:
            st.info("No hay cambios para el rango seleccionado.")
        else:
            expedientes_modificados = cambios_rango["clave_expediente"].dropna().unique().tolist()
            modificados = df[df["clave_expediente"].isin(expedientes_modificados)].copy()

            st.metric("Expedientes modificados", len(modificados))

            tabla = modificados.copy()
            tabla["⭐"] = tabla["Favorito"].apply(lambda x: "⭐" if x else "")

            st.dataframe(
                tabla[
                    [
                        "⭐",
                        "Nº Proced.",
                        "Nº Juzgado",
                        "Tipo órgano",
                        "Sección",
                        "Procedimiento",
                        "Materia",
                        "Fase Procesal",
                        "Último trámite",
                        "Fecha Últ. Trámite",
                        "Año Últ. Trámite",
                        "Estado F. Últ. Trámite",
                        "Nota",
                    ]
                ],
                use_container_width=True,
                height=450,
            )

            st.subheader("Cambios del rango")
            st.dataframe(
                cambios_rango.drop(columns=["clave_expediente"], errors="ignore"),
                use_container_width=True,
                height=350,
            )


with tab_detalle:
    st.subheader("Detalle de un expediente")

    if df.empty:
        st.info("Aún no hay expedientes.")
    else:
        opciones = (
            df["Nº Proced."].astype(str)
            + " | "
            + df["Procedimiento"].astype(str)
            + " | "
            + df["F. Aceptación"].astype(str)
        ).tolist()

        indice = st.selectbox("Selecciona expediente", range(len(opciones)), format_func=lambda i: opciones[i])
        seleccionado = df.iloc[indice]
        clave = seleccionado["clave_expediente"]

        detalle, cambios_detalle = cargar_detalle_expediente(clave)

        col1, col2 = st.columns([2, 1])

        with col1:
            st.markdown(f"### {detalle.get('numero_procedimiento', '')}")
            st.write(f"**Nº juzgado:** {detalle.get('juzgado_numero', '')}")
            st.write(f"**Tipo órgano:** {detalle.get('juzgado_tipo', '')}")
            st.write(f"**Sección:** {detalle.get('juzgado_seccion', '')}")
            st.write(f"**Órgano completo:** {detalle.get('organo_completo', detalle.get('juzgado', ''))}")
            st.write(f"**Procedimiento:** {detalle.get('procedimiento', '')}")
            st.write(f"**Materia:** {detalle.get('materia', '')}")
            st.write(f"**Fase procesal:** {detalle.get('fase_procesal', '')}")
            st.write(f"**Último trámite:** {detalle.get('ultimo_tramite', '')}")
            st.write(f"**Fecha último trámite:** {detalle.get('fecha_ultimo_tramite', '')}")
            st.write(f"**Primera importación:** {detalle.get('primera_importacion', '')}")
            st.write(f"**Última importación:** {detalle.get('ultima_importacion', '')}")

        with col2:
            favorito = st.checkbox(
                "Marcar como favorito",
                value=bool(detalle.get("favorito", 0)),
            )
            nota = st.text_area(
                "Nota interna",
                value=detalle.get("nota", "") or "",
                height=160,
            )

            if st.button("Guardar favorito/nota"):
                guardar_favorito_nota(clave, favorito, nota)
                st.success("Guardado. Recarga o cambia de pestaña para ver la tabla actualizada.")

        st.subheader("Histórico de este expediente")
        if cambios_detalle.empty:
            st.info("Este expediente todavía no tiene cambios históricos.")
        else:
            st.dataframe(
                cambios_detalle.drop(columns=["clave_expediente"], errors="ignore"),
                use_container_width=True,
                height=400,
            )

        with st.expander("Texto original extraído del PDF"):
            st.write(detalle.get("texto_original", ""))


with tab_cambios:
    st.subheader("Histórico completo de cambios")

    if df_cambios.empty:
        st.info("Aún no hay cambios registrados.")
    else:
        col1, col2 = st.columns(2)
        texto_cambio = col1.text_input("Buscar en cambios")
        campo = col2.multiselect(
            "Campo cambiado",
            sorted([x for x in df_cambios["Campo"].dropna().unique() if x]),
        )

        filtrado_cambios = df_cambios.copy()

        if campo:
            filtrado_cambios = filtrado_cambios[filtrado_cambios["Campo"].isin(campo)]

        filtrado_cambios = filtrar_por_texto(filtrado_cambios, texto_cambio)

        st.dataframe(
            filtrado_cambios.drop(columns=["clave_expediente"], errors="ignore"),
            use_container_width=True,
            height=650,
        )


with tab_importaciones:
    st.subheader("Importaciones realizadas")

    if df_importaciones.empty:
        st.info("Aún no hay importaciones registradas.")
    else:
        st.dataframe(
            df_importaciones,
            use_container_width=True,
            height=500,
        )


with tab_exportar:
    st.subheader("Exportar datos")

    df_filtrado_exportar = st.session_state.df_exportar_filtrado

    if df_filtrado_exportar is None or df_filtrado_exportar.empty:
        df_filtrado_exportar = df.copy()

    st.info(st.session_state.descripcion_exportacion)
    st.caption(
        "El Excel descargará exactamente la misma vista que aparece en la pestaña "
        "'Expedientes actuales': mismos filtros y mismas columnas visibles."
    )

    incluir_cambios = st.checkbox(
        "Incluir histórico de cambios completo",
        value=True,
        help="La hoja principal siempre será la vista filtrada. Esta opción añade el histórico completo como hoja adicional.",
    )

    incluir_importaciones = st.checkbox(
        "Incluir importaciones",
        value=True,
    )

    hojas = {
        "Vista filtrada": df_filtrado_exportar.drop(columns=["clave_expediente"], errors="ignore"),
    }

    if incluir_cambios:
        hojas["Cambios"] = df_cambios.drop(columns=["clave_expediente"], errors="ignore")

    if incluir_importaciones:
        hojas["Importaciones"] = df_importaciones

    excel = descargar_excel(hojas)

    nombre = f"expedientes_filtrados_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"

    st.download_button(
        "Descargar Excel filtrado",
        data=excel,
        file_name=nombre,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    st.info(
        "El Excel es útil para consultar y compartir datos, pero la copia completa recomendada "
        "es el backup SQLite de la barra lateral, porque conserva expedientes, históricos, favoritos y notas."
    )
